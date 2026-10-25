"""Загрузка вкладок FAQ из Google Sheets.

Правило воспроизводимости: пайплайн читает *снапшот* (CSV + meta.json в data/snapshots).
К живой таблице он обращается только если снапшота нет или указан refresh.

Скачивается вся таблица одним файлом xlsx и режется на вкладки по именам (gid не нужны),
каждая вкладка сохраняется отдельным CSV: `faq_<ключ>.csv` + `faq_<ключ>.meta.json`.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from ..config import FaqSheet


@dataclass
class RawFaqSheet:
    sheet: FaqSheet
    rows: list[list[str]]
    sha256: str
    path: Path
    snapshot_date: str          # YYYY-MM-DD
    fetched_at: str | None


def snapshot_csv_path(sheet: FaqSheet, snapshot_dir: Path) -> Path:
    return snapshot_dir / f"faq_{sheet.key}.csv"


def _meta_path(csv_path: Path) -> Path:
    return csv_path.with_suffix(".meta.json")


def _looks_like_html(data: bytes, content_type: str) -> bool:
    prefix = data[:200].lstrip().lower()
    return "text/html" in content_type.lower() or prefix.startswith((b"<!doctype html", b"<html"))


def fetch_workbook(sheet: FaqSheet, *, timeout: int = 180) -> bytes:
    """Скачивает всю таблицу как xlsx (нужен публичный доступ по ссылке)."""
    req = Request(sheet.xlsx_url, headers={"User-Agent": "Mozilla/5.0 BerserkIngest/0.1", "Accept": "*/*"})
    try:
        with urlopen(req, timeout=timeout) as resp:
            data = resp.read()
            ctype = resp.headers.get("content-type", "")
    except (HTTPError, URLError, TimeoutError) as exc:
        raise RuntimeError(f"Не удалось скачать таблицу {sheet.xlsx_url}: {type(exc).__name__}: {exc}") from exc
    if _looks_like_html(data, ctype):
        raise RuntimeError(f"{sheet.xlsx_url}: вместо xlsx пришёл HTML (нет публичного доступа по ссылке?)")
    return data


def _cell(value) -> str:
    """Значение ячейки xlsx -> строка так же, как её отдаёт CSV-экспорт Google (чекбокс = TRUE/FALSE)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def workbook_tabs(xlsx: bytes) -> dict[str, list[list[str]]]:
    """Все вкладки книги: имя -> строки (хвостовые пустые строки отброшены)."""
    from openpyxl import load_workbook   # ленивый импорт: нужен только для скачивания

    wb = load_workbook(io.BytesIO(xlsx), data_only=True)
    tabs: dict[str, list[list[str]]] = {}
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        while rows and not any(rows[-1]):
            rows.pop()
        tabs[ws.title.strip()] = rows
    return tabs


def rows_to_csv(rows: list[list[str]]) -> bytes:
    buf = io.StringIO(newline="")
    csv.writer(buf, lineterminator="\n").writerows(rows)
    return buf.getvalue().encode("utf-8")


def refresh_snapshots(sheets: list[FaqSheet], snapshot_dir: Path, *, xlsx: bytes | None = None) -> None:
    """Одна загрузка таблицы -> снапшоты всех запрошенных вкладок. xlsx можно передать готовый (для тестов)."""
    if not sheets:
        return
    print(f"[faq] скачиваю таблицу целиком для вкладок: {', '.join(s.set_name for s in sheets)} ...")
    data = xlsx if xlsx is not None else fetch_workbook(sheets[0])
    tabs = workbook_tabs(data)
    missing = [s.set_name for s in sheets if s.set_name not in tabs]
    if missing:
        raise RuntimeError(f"В таблице нет вкладок {missing}. Есть: {list(tabs)}")

    fetched_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    workbook_sha = hashlib.sha256(data).hexdigest()
    snapshot_dir.mkdir(parents=True, exist_ok=True)
    for sheet in sheets:
        body = rows_to_csv(tabs[sheet.set_name])
        path = snapshot_csv_path(sheet, snapshot_dir)
        path.write_bytes(body)
        _meta_path(path).write_text(json.dumps({
            "source_url": sheet.edit_url, "set_name": sheet.set_name, "format": "xlsx-export",
            "fetched_at": fetched_at, "sha256": hashlib.sha256(body).hexdigest(), "workbook_sha256": workbook_sha,
        }, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_csv(text: str) -> list[list[str]]:
    # StringIO(newline="") обязателен: в ячейках бывают переводы строк внутри кавычек,
    # а text.splitlines() рвёт такие ячейки пополам.
    rows = [list(r) for r in csv.reader(io.StringIO(text, newline=""))]
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def load_faq_sheet(
    sheet: FaqSheet,
    snapshot_dir: Path,
    *,
    refresh: bool = False,
    csv_path: Path | None = None,
) -> RawFaqSheet:
    path = Path(csv_path) if csv_path else snapshot_csv_path(sheet, snapshot_dir)
    meta_file = _meta_path(path)

    if refresh or not path.exists():
        if csv_path:
            if not path.exists():
                raise FileNotFoundError(f"Снапшот не найден: {path}")
        else:   # для нескольких вкладок сразу используйте refresh_snapshots (одна загрузка на всех)
            refresh_snapshots([sheet], snapshot_dir)

    data = path.read_bytes()
    sha = hashlib.sha256(data).hexdigest()

    fetched_at: str | None = None
    snapshot_date = None
    if meta_file.exists():
        meta = json.loads(meta_file.read_text(encoding="utf-8"))
        fetched_at = meta.get("fetched_at")
        if fetched_at:
            snapshot_date = fetched_at[:10]
    if snapshot_date is None:
        # без meta.json дата берётся из mtime файла и может «плыть» при git checkout
        snapshot_date = date.fromtimestamp(path.stat().st_mtime).isoformat()
        print(f"[faq] ВНИМАНИЕ: нет {meta_file.name}, snapshot_date взята из mtime ({snapshot_date})")

    rows = parse_csv(data.decode("utf-8-sig"))
    return RawFaqSheet(sheet, rows, sha, path, snapshot_date, fetched_at)
